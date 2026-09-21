# -*- coding: utf-8 -*-
"""
Deck toolkit for Aman's IP weekly updates.

Two jobs:
  1. Text measurement, so a box can be sized to its content BEFORE rendering.
     Overflow is the failure mode that makes a deck look shabby, and it is
     entirely preventable by measuring first.
  2. Shape helpers (boxes, chevrons, arrows, bars, title bar) in the locked
     house style, so every deck comes out looking like the same deck.

Fonts: measurement uses DejaVu because that is what LibreOffice substitutes
for Arial when rendering the check PDF, and DejaVu is wider than real Arial.
Measuring against the widest candidate means anything that fits in the check
render also fits in real PowerPoint. Do not "fix" this to Liberation — that
under-measures and text silently overflows.

Usage:
    import sys; sys.path.insert(0, '<path to this dir>')
    from deckkit import *
    prs = new_deck()
    s = slide(prs); titlebar(s, '1.  Where we left off')
    ...
    prs.save(out)
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import ImageFont
from lxml import etree

# ---------------------------------------------------------------- palette
NAVY    = RGBColor(0x22, 0x2A, 0x35)   # title bar, dark bands, "conclusion" boxes
SLATE   = RGBColor(0x3D, 0x4B, 0x5C)
STEEL   = RGBColor(0x2F, 0x5D, 0x8A)   # primary accent (eyebrows, layer 3)
STEEL_L = RGBColor(0xE4, 0xEC, 0xF4)
RUST    = RGBColor(0xB5, 0x4F, 0x35)   # "this is the important/critical one"
RUST_L  = RGBColor(0xFA, 0xEC, 0xE8)
TEAL    = RGBColor(0x2D, 0x6E, 0x63)
TEAL_L  = RGBColor(0xE4, 0xEF, 0xED)
LIGHT   = RGBColor(0xF2, 0xF4, 0xF7)   # default card fill
BORDER  = RGBColor(0xC8, 0xD0, 0xD8)
INK     = RGBColor(0x1A, 0x1A, 0x1A)
MUTED   = RGBColor(0x5A, 0x65, 0x70)   # explanatory prose
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
ONDARK  = RGBColor(0xC8, 0xD2, 0xDC)   # body text on NAVY
ARIAL   = "Arial"

# ---------------------------------------------------------------- grid
SW, SH       = 12192000, 6858000       # 16:9
MX           = 502920                  # side margin
CW           = SW - 2 * MX             # content width
BAR_Y, BAR_H = 275529, 461665          # title bar
BODY_TOP     = 950000                  # first element under the title bar
BODY_BOTTOM  = 6500000                 # keep everything above this

# ---------------------------------------------------------------- measuring
_FP = {('sans', False): '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
       ('sans', True):  '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
       ('serif', False):'/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf',
       ('serif', True): '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'}
_cache = {}
SAFETY = 0.98

def _font(pt, bold, fam='sans'):
    k = (round(pt * 4), bold, fam)
    if k not in _cache:
        _cache[k] = ImageFont.truetype(_FP[(fam, bold)], max(1, int(round(pt * 4))))
    return _cache[k]

def text_w_pt(s, pt, bold=False, fam='sans'):
    return _font(pt, bold, fam).getlength(s) / 4.0

def n_lines(text, pt, width_emu, bold=False, fam='sans'):
    """Greedy word wrap, the same rule PowerPoint uses. Returns a line count."""
    limit = width_emu / 12700.0 * SAFETY
    lines, cur = 0, ''
    for w in text.split():
        trial = w if not cur else cur + ' ' + w
        if text_w_pt(trial, pt, bold, fam) <= limit or not cur:
            cur = trial
        else:
            lines += 1; cur = w
    return lines + (1 if cur else 0)

def text_h(text, pt, width_emu, line_pt=None, bold=False, fam='sans'):
    """Height in EMU that `text` needs at `pt` inside `width_emu`."""
    return int(n_lines(text, pt, width_emu, bold, fam) * (line_pt or pt * 1.4) * 12700)

# ---------------------------------------------------------------- primitives
def new_deck():
    p = Presentation()
    p.slide_width, p.slide_height = Emu(SW), Emu(SH)
    return p

def slide(prs, bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])      # blank
    if bg is not None:
        xml = ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               '<p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
               '<a:effectLst/></p:bgPr></p:bg>' % str(bg))
        s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld'
                        ).insert(0, etree.fromstring(xml))
    return s

def _fill(sh, fill, line, lw=1.0):
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)

def _settext(sh, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, pad=90000):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(pad)
    tf.margin_top = tf.margin_bottom = Emu(40000)
    tf.vertical_anchor = anchor
    for i, sp in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = sp.get('align', align)
        if sp.get('sb'): p.space_before = Pt(sp['sb'])
        p.line_spacing = Pt(sp.get('size', 12) * sp.get('lh', 1.25))
        r = p.add_run(); r.text = sp['t']
        f = r.font
        f.name = ARIAL; f.size = Pt(sp.get('size', 12))
        f.bold = sp.get('bold', False); f.italic = sp.get('italic', False)
        f.color.rgb = sp.get('color', INK)

def box(s, x, y, w, h, lines, fill=LIGHT, line=BORDER, lw=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, pad=90000):
    sh = s.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sh.adjustments[0] = adj
        except Exception: pass
    _fill(sh, fill, line, lw); _settext(sh, lines, align, anchor, pad)
    return sh

def chev(s, x, y, w, h, lines, fill=LIGHT, line=BORDER, lw=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    _fill(sh, fill, line, lw); _settext(sh, lines)
    return sh

def arrow(s, x, y, w, h, direction='right', fill=BORDER):
    shp = {'right': MSO_SHAPE.RIGHT_ARROW, 'down': MSO_SHAPE.DOWN_ARROW}[direction]
    sh = s.shapes.add_shape(shp, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    _fill(sh, fill, None)
    return sh

def txt(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, sp in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = sp.get('align', align)
        if sp.get('sb'): p.space_before = Pt(sp['sb'])
        p.line_spacing = Pt(sp.get('size', 12) * sp.get('lh', 1.3))
        r = p.add_run(); r.text = sp['t']
        f = r.font; f.name = ARIAL; f.size = Pt(sp.get('size', 12))
        f.bold = sp.get('bold', False); f.italic = sp.get('italic', False)
        f.color.rgb = sp.get('color', INK)
    return tb

# ---------------------------------------------------------------- house parts
def titlebar(s, title):
    """Full-bleed dark bar. Number every slide ('3.  ...') so the reader can cite one."""
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(BAR_Y), Emu(SW), Emu(BAR_H))
    _fill(sh, NAVY, None)
    _settext(sh, [{'t': title, 'size': 22, 'bold': True, 'color': WHITE}])
    return sh

def eyebrow(s, x, y, w, text, color=STEEL, size=11):
    """Small caps-ish label above a figure. Make it say the point, not the category."""
    return txt(s, x, y, w, 220000, [{'t': text, 'size': size, 'bold': True, 'color': color}])

def caption(s, x, y, w, text, size=10, color=MUTED, align=PP_ALIGN.LEFT, lh=1.35):
    """Auto-height explanatory prose. Returns the y where the next element can start."""
    h = text_h(text, size, w, size * lh)
    txt(s, x, y, w, h, [{'t': text, 'size': size, 'color': color, 'lh': lh, 'align': align}])
    return y + h

def band(s, y, label, body, dark=False, size=12.5, line=18.5, pad=240000):
    """Auto-height full-width callout. Use NAVY (dark=True) for the one conclusion
    per slide you want a skimming reader to take away. Returns the bottom y."""
    iw = CW - 2 * pad
    bh = text_h(body, size, iw, line)
    h = 150000 + 250000 + 60000 + bh + 150000
    box(s, MX, y, CW, h, [], fill=NAVY if dark else LIGHT,
        line=None if dark else BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, MX + pad, y + 150000, iw, 250000,
        [{'t': label, 'size': 10.5, 'bold': True, 'color': WHITE if dark else STEEL}])
    txt(s, MX + pad, y + 460000, iw, bh,
        [{'t': body, 'size': size, 'color': ONDARK if dark else INK, 'lh': line / size}])
    return y + h

def barchart(s, x, y, rows, track=760000, row_h=530000, name_w=1250000,
             q_w=3050000, label_size=11):
    """Weight/score bars with a name, the question it asks, and why it matters.
    rows = [(name, value, color, question, why), ...]; value scales against the max.
    Returns the bottom y. A bar alone means nothing to a cold reader — the
    'why' column is the part that does the work."""
    top = max(r[1] for r in rows)
    for name, val, col, q, wh in rows:
        box(s, x, y + 60000, track * val / float(top), 130000, [], fill=col,
            line=None, shape=MSO_SHAPE.RECTANGLE)
        txt(s, x + track + 130000, y - 25000, name_w, 240000,
            [{'t': '%s  %s' % (name, val), 'size': label_size, 'bold': True, 'color': NAVY}])
        txt(s, x + track + name_w + 250000, y - 25000, q_w, 400000,
            [{'t': q, 'size': 10.5, 'color': INK, 'lh': 1.3}])
        ex = x + track + name_w + q_w + 400000
        txt(s, ex, y - 25000, (x + CW) - ex, 500000,
            [{'t': wh, 'size': 9.5, 'color': MUTED, 'lh': 1.3}])
        y += row_h
    return y

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text
