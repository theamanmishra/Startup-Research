#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Check a .pptx for overflow, then render every slide to PNG so it can be looked at.

Why this exists: the only reliable way to know a deck is not shabby is to look
at it. Text that overflows its box, or runs off the slide, is invisible in the
code and obvious to the person receiving the deck. This does both passes:

  1. Static check — for every shape, measure the text against the box it sits
     in and against the slide edges. Catches most problems in a second, with
     no rendering.
  2. Visual render — PDF via LibreOffice, then PNG per slide, so the agent can
     Read the images and catch what geometry cannot (collisions between
     separate shapes, visual crowding, a figure that reads wrong).

Usage:
    python3 render_check.py deck.pptx                 # check + render
    python3 render_check.py deck.pptx --check-only    # skip rendering
    python3 render_check.py deck.pptx --outdir /tmp/r --scale 2

First run on a fresh box may need:
    apt-get update -q && apt-get install -y -q libreoffice-impress
    pip install python-pptx pypdfium2 pillow
LibreOffice ships without Impress in some images; without it, every convert
fails with the unhelpful "source file could not be loaded".
"""
import argparse, os, subprocess, sys, glob

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import n_lines, SH, SW, BODY_BOTTOM          # noqa: E402
from pptx import Presentation                              # noqa: E402
from pptx.util import Emu                                  # noqa: E402

EMU_IN = 914400.0


def static_check(path, bottom_limit=BODY_BOTTOM, tol=0.06):
    """Measure each shape's text against its own box and the slide edges.

    The two cases behave differently and must be judged differently:

      * A filled or bordered card (an autoshape) shows its overflow — the text
        crosses the border and looks broken. Flag it.
      * A plain textbox has no visible edge, so text running past its nominal
        height just continues into whitespace. Harmless. Only worth flagging
        if it runs off the slide or below the safe bottom margin.

    `tol` (inches) absorbs the deliberate over-measurement in deckkit: DejaVu
    is wider than the Arial PowerPoint will actually use, so a hair of
    predicted overflow usually renders fine. Anything past the tolerance is
    worth a look.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    problems = []
    for idx, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame or sh.width is None or sh.height is None:
                continue
            tf = sh.text_frame
            txt = ' '.join(tf.text.split())
            ml = tf.margin_left or 0
            mr = tf.margin_right or 0
            inner = max(1, sh.width - ml - mr)

            need = 0
            for p in tf.paragraphs:
                runs = [r for r in p.runs if r.text.strip()]
                if not runs:
                    continue
                size = next((r.font.size.pt for r in runs if r.font.size), 12)
                bold = bool(runs[0].font.bold)
                line = p.line_spacing.pt if hasattr(p.line_spacing, 'pt') else size * 1.3
                sb = p.space_before.pt if hasattr(p.space_before, 'pt') else 0
                need += (sb + n_lines(''.join(r.text for r in runs), size, inner, bold) * line) * 12700

            is_card = sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX
            over = need - sh.height
            if txt and is_card and over > tol * EMU_IN:
                problems.append((idx, 'text overflows its card by %.2f in' % (over / EMU_IN), txt[:64]))

            text_bottom = sh.top + (need if not is_card else sh.height)
            if sh.top + sh.height > SH or text_bottom > SH:
                problems.append((idx, 'runs off the bottom of the slide', txt[:64]))
            elif txt and text_bottom > bottom_limit + tol * EMU_IN:
                problems.append((idx, 'sits %.2f in below the safe bottom margin'
                                 % ((text_bottom - bottom_limit) / EMU_IN), txt[:64]))
            if sh.left < 0 or sh.left + sh.width > SW:
                problems.append((idx, 'extends past the left or right edge', txt[:64]))
    return len(prs.slides._sldIdLst), problems


def render(path, outdir, scale=1.5):
    os.makedirs(outdir, exist_ok=True)
    profile = os.path.join(outdir, '_loprofile')
    cmd = ['soffice', '-env:UserInstallation=file://' + profile, '--headless',
           '--norestore', '--convert-to', 'pdf', '--outdir', outdir, path]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + '.pdf')
    if not os.path.exists(pdf):
        print('RENDER FAILED. soffice said:\n' + (r.stdout or '') + (r.stderr or ''))
        print('If it says "source file could not be loaded", Impress is missing:')
        print('  apt-get update -q && apt-get install -y -q libreoffice-impress')
        return []
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf)
    pages = []
    for i in range(len(doc)):
        p = os.path.join(outdir, 'slide%02d.png' % (i + 1))
        doc[i].render(scale=scale).to_pil().save(p)
        pages.append(p)
    return pages


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx')
    ap.add_argument('--outdir', default=None)
    ap.add_argument('--scale', type=float, default=1.5)
    ap.add_argument('--check-only', action='store_true')
    a = ap.parse_args()

    outdir = a.outdir or os.path.join(os.path.dirname(os.path.abspath(a.pptx)), 'render')
    n, problems = static_check(a.pptx)
    print('%d slides checked.' % n)
    if problems:
        print('\n%d possible problem(s):' % len(problems))
        for slide_no, what, snippet in problems:
            print('  slide %-2d  %-46s  %s' % (slide_no, what, snippet))
    else:
        print('No geometric overflow found.')

    if not a.check_only:
        pages = render(a.pptx, outdir, a.scale)
        if pages:
            print('\nRendered %d PNG(s) to %s' % (len(pages), outdir))
            print('Now LOOK at them — Read each PNG. Geometry cannot see a collision '
                  'between two separate shapes, or a figure that simply reads wrong.')
    return 1 if problems else 0


if __name__ == '__main__':
    sys.exit(main())
